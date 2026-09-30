"""Parse the odd-semester Master and sectionwise academic workbooks."""

from collections import defaultdict
from hashlib import sha1
from io import BytesIO
import re

from openpyxl import load_workbook


PERIOD_TIMES = {
    1: ("09:30 AM", "10:20 AM"),
    2: ("10:20 AM", "11:10 AM"),
    3: ("11:25 AM", "12:15 PM"),
    4: ("12:15 PM", "01:05 PM"),
    6: ("01:45 PM", "02:35 PM"),
    7: ("02:35 PM", "03:25 PM"),
    8: ("03:25 PM", "04:15 PM"),
}

DEPARTMENTS = {
    "CSE": "Computer Science and Engineering",
    "AI": "Artificial Intelligence and Machine Learning",
    "AIML": "Artificial Intelligence and Machine Learning",
    "AI&ML": "Artificial Intelligence and Machine Learning",
    "AIDS": "AI and Data Science",
    "AI&DS": "AI and Data Science",
    "ECE": "Electronics and Communication Engineering",
    "ME": "Mechanical Engineering",
}

DAYS = {
    "M": "Mon", "MON": "Mon",
    "T": "Tue", "TU": "Tue", "TUE": "Tue",
    "W": "Wed", "WED": "Wed",
    "TH": "Thu", "THU": "Thu",
    "F": "Fri", "FRI": "Fri",
}


def _academic_year_number(value):
    text = str(value or "").casefold()
    if "second year" in text or "year: ii" in text:
        return 2
    if "third year" in text or "year: iii" in text:
        return 3
    if "fourth year" in text or "final year" in text or "year: iv" in text:
        return 4
    return None


def _room_key(value):
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def _parse_departments(value):
    result = []
    for token in re.split(r"[,/]", str(value or "")):
        key = " ".join(token.upper().split()).strip(" .:")
        name = DEPARTMENTS.get(key)
        if name and name not in result:
            result.append(name)
    return result


def _read_section_metadata(workbook_bytes):
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    section_map = defaultdict(list)
    warnings = []

    for sheet in workbook.worksheets:
        sheet_year = _academic_year_number(sheet.title.replace("_", " "))
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if not isinstance(value, str) or "section" not in value.casefold():
                    continue
                match = re.search(r"Section\s*:\s*([A-J])(?=\s|\()", value, re.IGNORECASE)
                if not match or "B.SC" in value.upper() or "BCA" in value.upper():
                    continue

                year_match = re.search(r"Year\s*:?\s*(IV|III|II|I)\b", value, re.IGNORECASE)
                roman_year = {"I": 1, "II": 2, "III": 3, "IV": 4}
                year = roman_year.get(year_match.group(1).upper()) if year_match else sheet_year
                if year not in (2, 3, 4):
                    continue

                department_match = re.search(r"Dept\.?\s*:?\s*([^\n]+)", value, re.IGNORECASE)
                departments = _parse_departments(department_match.group(1)) if department_match else []
                room_match = re.search(
                    r"\(\s*((?:R|G)-?\s*\d+[A-Z]?)\s*\)",
                    value,
                    re.IGNORECASE
                ) or re.search(
                    r"\b((?:R|G)-?\s*\d+[A-Z]?)\b",
                    value.split("Year", 1)[0],
                    re.IGNORECASE
                )
                if not departments or not room_match:
                    warnings.append(f"Unmatched section header: {value.strip()}")
                    continue

                section = match.group(1).upper()
                room = room_match.group(1).replace(" ", "")
                entry = {"departments": departments, "room": room}
                key = (year, section, _room_key(room))
                if entry not in section_map[key]:
                    section_map[key].append(entry)

    return section_map, warnings


def _parse_days(value):
    tokens = [token.upper() for token in re.split(r"[,;/\s]+", str(value or "").strip()) if token]
    return [DAYS[token] for token in tokens if token in DAYS]


def _parse_periods(value):
    periods = []
    for token in re.split(r"[,;/\s]+", str(value or "").strip()):
        if not token:
            continue
        if re.fullmatch(r"\d+\s*-\s*\d+", token):
            start, end = map(int, re.split(r"\s*-\s*", token))
            periods.extend(range(start, end + 1))
        elif token.isdigit():
            periods.append(int(token))
    return periods


def parse_academic_workbooks(master_bytes, sectionwise_bytes, academic_batch):
    section_map, warnings = _read_section_metadata(sectionwise_bytes)
    workbook = load_workbook(BytesIO(master_bytes), read_only=True, data_only=True)
    courses = {}
    structures = {}
    sections = {}
    timetable = []
    skipped = 0
    current_year = None
    current_course = None

    for row_number, row in enumerate(workbook.active.iter_rows(values_only=True), 1):
        if isinstance(row[0], str) and "year" in row[0].casefold():
            group_header = row[0].upper()
            is_btech_group = "B.TECH" in group_header or "BTECH" in group_header
            is_other_program = "BCA" in group_header or "BSC" in group_header
            current_year = (
                _academic_year_number(row[0])
                if is_btech_group and not is_other_program
                else None
            )
            current_course = None
            continue

        if row[1] and row[2] and str(row[1]).strip().casefold() != "code":
            current_year = current_year or _academic_year_number(row[2])
            if current_year not in (2, 3, 4):
                current_course = None
                continue
            load_pattern = [int(value) for value in re.findall(r"\d+", str(row[3] or ""))]
            current_course = {
                "source_code": str(row[1]).strip(),
                "name": str(row[2]).strip(),
                "credits": load_pattern[-1] if load_pattern else 0,
                "has_lab": len(load_pattern) >= 3 and load_pattern[2] > 0,
                "year": current_year,
                "semester": current_year * 2 - 1,
            }

        section_value = str(row[4] or "").strip()
        section_match = re.match(r"Section\s*[- ]*([A-J])(?=\s|\(|$)", section_value, re.IGNORECASE)
        if not section_match or current_course is None:
            continue

        section = section_match.group(1).upper()
        room = str(row[5] or "").strip()
        day_tokens = _parse_days(row[6])
        period_tokens = _parse_periods(row[7])
        faculty_name = str(row[8] or "").strip()
        if not faculty_name or not day_tokens or len(day_tokens) != len(period_tokens):
            skipped += 1
            continue

        key = (current_year, section, _room_key(room))
        metadata = section_map.get(key)
        if not metadata:
            candidates = [
                values for (year, section_name, _), values in section_map.items()
                if year == current_year and section_name == section
            ]
            metadata = candidates[0] if len(candidates) == 1 else None
        if not metadata:
            skipped += 1
            continue

        for section_data in metadata:
            for department in section_data["departments"]:
                semester = current_course["semester"]
                structure_key = (department, current_year, semester, str(academic_batch))
                structures[structure_key] = {
                    "department": department,
                    "year": current_year,
                    "semester": semester,
                    "academic_batch": str(academic_batch),
                }
                digest = sha1(
                    f"{current_course['source_code']}|{current_course['name']}|{department}|{current_year}|{semester}".encode()
                ).hexdigest()[:10].upper()
                internal_code = f"IMP-{digest}"
                courses[internal_code] = {
                    "internal_code": internal_code,
                    "catalog_code": current_course["source_code"],
                    "course_name": current_course["name"],
                    "department": department,
                    "year": current_year,
                    "semester": semester,
                    "academic_batch": str(academic_batch),
                    "credits": current_course["credits"],
                    "course_type": "Theory+Lab" if current_course["has_lab"] else "Theory",
                    "structure_key": structure_key,
                }
                sections[(department, current_year, semester, section)] = {
                    "department": department,
                    "year": current_year,
                    "semester": semester,
                    "section_name": section,
                    "room_number": section_data["room"],
                }
                for day, period in zip(day_tokens, period_tokens):
                    period_time = PERIOD_TIMES.get(period)
                    if period_time is None:
                        skipped += 1
                        continue
                    timetable.append({
                        "department": department,
                        "year": current_year,
                        "semester": semester,
                        "academic_batch": str(academic_batch),
                        "section": section,
                        "day": day,
                        "slot": f"P{period}",
                        "start_time": period_time[0],
                        "end_time": period_time[1],
                        "course_code": current_course["source_code"],
                        "course_name": current_course["name"],
                        "faculty_name": faculty_name,
                        "room_number": room or section_data["room"],
                        "internal_course_code": internal_code,
                    })

    timetable_groups = defaultdict(list)
    for entry in timetable:
        group_key = (
            entry["department"],
            entry["year"],
            entry["semester"],
            entry["section"],
            entry["day"],
            entry["slot"],
        )
        timetable_groups[group_key].append(entry)

    conflict_keys = set()
    conflicts = []
    for group_key, entries in timetable_groups.items():
        distinct_courses = {entry["internal_course_code"] for entry in entries}
        if len(distinct_courses) > 1:
            conflict_keys.add(group_key)
            conflicts.append({
                "department": group_key[0],
                "year": group_key[1],
                "semester": group_key[2],
                "section": group_key[3],
                "day": group_key[4],
                "slot": group_key[5],
                "courses": ", ".join(sorted({entry["course_code"] for entry in entries})),
            })

    if conflicts:
        warnings.append(
            f"Skipped {len(conflicts)} timetable periods with multiple courses assigned to the same section."
        )
    if skipped:
        warnings.append(f"Skipped {skipped} unsupported or unmatched section/period entries.")
    return {
        "courses": list(courses.values()),
        "structures": list(structures.values()),
        "sections": list(sections.values()),
        "timetable": [
            entry
            for entry in timetable
            if (
                entry["department"],
                entry["year"],
                entry["semester"],
                entry["section"],
                entry["day"],
                entry["slot"],
            ) not in conflict_keys
        ],
        "conflicts": conflicts,
        "warnings": warnings,
    }
